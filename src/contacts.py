"""Contact-file loading for the call and messaging scripts."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Contact:
    number: str
    name: str = ""


def _as_text(value: object) -> str:
    """Convert a spreadsheet cell to text without a trailing `.0`."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _contacts_from_rows(rows: list[dict[str, object]], source: Path) -> list[Contact]:
    contacts = []
    for row_number, row in enumerate(rows, start=2):
        normalized = {str(key).strip().lower(): value for key, value in row.items() if key}
        number = _as_text(normalized.get("number"))
        name = _as_text(normalized.get("name"))
        if not number:
            if any(_as_text(value) for value in normalized.values()):
                raise ValueError(f"Missing number in {source} at row {row_number}.")
            continue
        contacts.append(Contact(number=number, name=name))
    if not contacts:
        raise ValueError(f"No contacts found in {source}. Add a 'number' column and at least one row.")
    return contacts


def load_contacts(path: Path) -> list[Contact]:
    """Load `number` and optional `name` columns from a CSV or XLSX contact file."""
    if not path.exists():
        raise FileNotFoundError(
            f"Contact file not found: {path}. Copy data/contacts.example.csv to data/contacts.csv "
            "and add your contacts."
        )

    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as file:
            reader = csv.DictReader(file)
            if not reader.fieldnames or "number" not in {
                field.strip().lower() for field in reader.fieldnames
            }:
                raise ValueError(f"{path} must have a 'number' column.")
            return _contacts_from_rows(list(reader), path)

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError(
                "Excel files require openpyxl. Install it with: python3 -m pip install openpyxl"
            ) from error
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError(f"{path} is empty.")
        fieldnames = [_as_text(header) for header in headers]
        if "number" not in {field.strip().lower() for field in fieldnames}:
            raise ValueError(f"{path} must have a 'number' column.")
        records = [dict(zip(fieldnames, row)) for row in rows]
        return _contacts_from_rows(records, path)

    raise ValueError(f"Unsupported contact file {path}. Use .csv or .xlsx.")
