import csv, os, re, subprocess, sys, time

VERSION = "whatsapp_call v4"

# ---------------- SETTINGS ----------------
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # project folder (parent of scripts/)
CONTACTS_FILE = os.path.join(ROOT, "data", "contacts.csv")           # .csv or .xlsx

AUDIO = "voice_f.mp3"    # file name only; it is searched for inside the project folder
PLAY_AUDIO = True        # False = just stay on the call for MAX_PLAY seconds
RING_TIMEOUT = 45        # max seconds to wait for the person to answer
MAX_PLAY = 30            # call/audio duration after answer
GAP_BETWEEN_CALLS = 5    # seconds between calls
PKG = "com.whatsapp"     # use "com.whatsapp.w4b" for WhatsApp Business

# Button labels on your phone (found with:  python3 scripts/whatsapp_call.py labels)
VOICE_CALL_BTN = r"^voice call$"
IN_CALL_BTN = r"^(leave call|end call|hang up|disconnect)$"
TIMER = r"\b\d{1,2}:\d{2}(:\d{2})?\b"     # call timer (like 00:05) appears after answer
DEBUG = True             # print every new label seen on the call screen while waiting

# The call screen is detected by the name of the focused window, not by button labels.
# Check it during a call with:  python3 scripts/whatsapp_call.py focus
CALL_SCREEN = r"voip"

# How to cut the call
END_TAP = (840, 1955)            # position of the red "Leave call" button on your phone
FORCE_STOP_AFTER_HANGUP = True   # guarantee: force-stop WhatsApp after hang up (re-opened at the end)

# If the timer is never detected, treat the call as answered after this many seconds
# while the call screen is still open. None = only trust the timer.
ASSUME_ANSWERED_AFTER = None
# ------------------------------------------


# ---------- contacts (CSV / Excel) ----------
NUMBER_COLS = ("number", "phone", "mobile", "whatsapp", "contact", "cell", "phone number")
NAME_COLS = ("name", "full name", "contact name")


def cell_to_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                    # Excel stores numbers as floats: 3359863734.0
    return str(v).strip()


def normalize(value):
    """0336-1234567, +92 336 1234567, 3361234567 -> 923361234567"""
    digits = re.sub(r"\D", "", cell_to_str(value))
    if digits.startswith("00"):
        digits = digits[2:]
    if digits.startswith("0"):
        digits = "92" + digits[1:]
    elif len(digits) == 10 and digits.startswith("3"):
        digits = "92" + digits         # Excel dropped the leading zero
    return digits


def load_contacts(path):
    """Returns a list of (number, name). Replace this with your own loader if you like."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook          # pip install openpyxl
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    rows = [r for r in rows if any(cell_to_str(c) for c in r)]
    if not rows:
        return []

    header = [cell_to_str(c).lower() for c in rows[0]]
    num_i = next((i for i, h in enumerate(header) if h in NUMBER_COLS), None)
    name_i = next((i for i, h in enumerate(header) if h in NAME_COLS), None)
    data = rows[1:]
    if num_i is None:                  # no header row: first column is the number
        num_i = 0
        if re.sub(r"\D", "", header[0]):
            data = rows                # first row is already a number

    contacts = []
    for r in data:
        number = normalize(r[num_i]) if num_i < len(r) else ""
        if not (10 <= len(number) <= 15):
            print(f"  skipping invalid number: {r}")
            continue
        name = cell_to_str(r[name_i]) if name_i is not None and name_i < len(r) else ""
        contacts.append((number, name))
    return contacts


def find_audio():
    for dirpath, dirnames, files in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if d not in (".git", "node_modules", "__pycache__")]
        if AUDIO in files:
            return os.path.join(dirpath, AUDIO)
    return None


# ---------- phone helpers ----------
def adb(*args):
    return subprocess.run(["adb", "shell", *args],
                          capture_output=True, text=True).stdout


LAST_DUMP_MSG = ""


def ui_nodes():
    """Dump the current screen and return its elements with tap coordinates."""
    global LAST_DUMP_MSG
    adb("rm", "-f", "/sdcard/ui.xml")            # avoid reading a stale dump
    r = subprocess.run(["adb", "shell", "uiautomator", "dump", "/sdcard/ui.xml"],
                       capture_output=True, text=True)
    LAST_DUMP_MSG = (r.stdout + r.stderr).strip()
    xml = adb("cat", "/sdcard/ui.xml")
    nodes = []
    for node in re.findall(r"<node [^>]*>", xml):
        text = re.search(r' text="([^"]*)"', node)
        desc = re.search(r'content-desc="([^"]*)"', node)
        b = re.search(r'bounds="\[(\d+),(\d+)\]\[(\d+),(\d+)\]"', node)
        if not b:
            continue
        x1, y1, x2, y2 = map(int, b.groups())
        nodes.append({
            "text": text.group(1) if text else "",
            "desc": desc.group(1) if desc else "",
            "x": (x1 + x2) // 2,
            "y": (y1 + y2) // 2,
        })
    return nodes


def find(nodes, pattern):
    for n in nodes:
        if re.search(pattern, n["text"].strip(), re.I) or \
           re.search(pattern, n["desc"].strip(), re.I):
            return n
    return None


def tap(n):
    adb("input", "tap", str(n["x"]), str(n["y"]))


def tap_label(pattern, tries=6, delay=1.5):
    for i in range(tries):
        n = find(ui_nodes(), pattern)
        if n:
            tap(n)
            return True
        if i < tries - 1:
            time.sleep(delay)
    return False


def focus_line():
    """Which window is on screen right now (works even when uiautomator does not)."""
    return adb("dumpsys window | grep -E 'mCurrentFocus|mFocusedApp'").strip()


def in_call_screen():
    return bool(re.search(CALL_SCREEN, focus_line(), re.I))


# ---------- call logic ----------
def wait_for_answer():
    """Returns 'answered', 'answered (assumed)', 'failed' or 'timeout'."""
    start = time.time()
    seen = set()
    while time.time() - start < RING_TIMEOUT:
        elapsed = time.time() - start
        nodes = ui_nodes()
        if DEBUG:
            for n in nodes:
                for label in (n["text"], n["desc"]):
                    if label and label not in seen:
                        seen.add(label)
                        print(f"    screen: {label}")
        # timer text anywhere below the status bar means the call is connected
        for n in nodes:
            if n["y"] > 120 and (re.search(TIMER, n["text"]) or
                                 re.search(TIMER, n["desc"])):
                return "answered"
        if elapsed > 6 and not in_call_screen():
            return "failed"                  # call screen closed: declined / not on WhatsApp
        if ASSUME_ANSWERED_AFTER and elapsed >= ASSUME_ANSWERED_AFTER:
            return "answered (assumed)"
        time.sleep(1)
    return "timeout"


def play_audio(audio_path):
    player = None
    if PLAY_AUDIO and audio_path:
        player = subprocess.Popen(["ffplay", "-nodisp", "-autoexit", "-loglevel",
                                   "quiet", "-t", str(MAX_PLAY), audio_path])
    elif PLAY_AUDIO:
        print(f"  audio file '{AUDIO}' not found, staying on the call without audio")
    t0 = last = time.time()
    while time.time() - t0 < MAX_PLAY:
        time.sleep(1)
        if player and player.poll() is not None:
            break                            # audio finished
        if time.time() - last > 4:
            last = time.time()
            if not in_call_screen():
                print("  call ended early")
                break
    if player and player.poll() is None:
        player.terminate()


def hang_up():
    if tap_label(IN_CALL_BTN, tries=1, delay=0):
        print("  hang up: tapped the button")
    elif END_TAP and in_call_screen():
        adb("input", "tap", str(END_TAP[0]), str(END_TAP[1]))
        print("  hang up: tapped the fixed position")
    else:
        print("  hang up: no call screen found")
    time.sleep(2)
    if FORCE_STOP_AFTER_HANGUP:
        adb("am", "force-stop", PKG)         # guarantees the call is dropped
        time.sleep(1)


def main():
    print(VERSION)

    # helper modes for finding the right settings
    if len(sys.argv) > 1 and sys.argv[1] == "labels":
        found = False
        for n in ui_nodes():
            if n["text"] or n["desc"]:
                found = True
                print(f'text="{n["text"]}"  desc="{n["desc"]}"  at ({n["x"]},{n["y"]})')
        if not found:
            print("No labels found on this screen.")
            print("Dump message:", LAST_DUMP_MSG or "(empty)")
        return
    if len(sys.argv) > 1 and sys.argv[1] == "focus":
        print(focus_line() or "(empty)")
        print("call screen detected:", in_call_screen())
        return

    state = subprocess.run(["adb", "get-state"], capture_output=True, text=True)
    if "device" not in state.stdout:
        print("Phone not connected. Run 'adb devices' and allow USB debugging.")
        sys.exit(1)

    if not os.path.exists(CONTACTS_FILE):
        print(f"Contacts file not found: {CONTACTS_FILE}")
        sys.exit(1)
    contacts = load_contacts(CONTACTS_FILE)
    if not contacts:
        print("No valid numbers found in the contacts file.")
        sys.exit(1)
    print(f"Loaded {len(contacts)} contacts from {os.path.basename(CONTACTS_FILE)}")

    audio_path = find_audio() if PLAY_AUDIO else None

    adb("input", "keyevent", "KEYCODE_WAKEUP")
    adb("svc", "power", "stayon", "usb")      # keep the screen on while plugged in

    results = []
    try:
        for i, (number, name) in enumerate(contacts, 1):
            who = f"{name} ({number})" if name else number
            print(f"[{i}/{len(contacts)}] WhatsApp calling {who}...")
            adb("am", "start", "-a", "android.intent.action.VIEW",
                "-d", f"'whatsapp://send?phone={number}'", "-p", PKG)
            time.sleep(3)

            if not tap_label(VOICE_CALL_BTN):
                print("  could not find the Voice call button (number not on WhatsApp?)")
                results.append((number, "no call button"))
                adb("input", "keyevent", "KEYCODE_BACK")
                continue

            time.sleep(1)
            confirm = find(ui_nodes(), r"^call$")     # confirmation dialog, if any
            if confirm:
                tap(confirm)

            status = wait_for_answer()
            if status.startswith("answered"):
                print(f"  {status}, playing audio")
                play_audio(audio_path)
            else:
                print(f"  {status}, skipping")
            results.append((number, status))

            hang_up()
            time.sleep(GAP_BETWEEN_CALLS)

    except KeyboardInterrupt:
        print("\nStopped by user.")
        hang_up()

    finally:
        if FORCE_STOP_AFTER_HANGUP:           # open WhatsApp again so it keeps receiving messages
            adb("monkey", "-p", PKG, "-c", "android.intent.category.LAUNCHER", "1")

    print("\nSummary:")
    for number, r in results:
        print(f"  {number}: {r}")
    print("All done.")


if __name__ == "__main__":
    main()
    sys.exit(0)